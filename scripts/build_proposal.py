from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(__file__).resolve().parents[1]
(root / "docs/proposal").mkdir(parents=True, exist_ok=True)
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin, sec.bottom_margin = Inches(0.75), Inches(0.75)
sec.left_margin = sec.right_margin = Inches(0.85)
for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2"]:
    style = doc.styles[name]
    style.font.name = "Arial"
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.size = Pt(11)
normal = doc.styles["Normal"].paragraph_format
normal.line_spacing = 1.12
normal.space_after = Pt(9)
doc.styles["Title"].font.size = Pt(25)
doc.styles["Title"].font.bold = True
doc.styles["Title"].paragraph_format.space_after = Pt(12)
doc.styles["Heading 1"].font.size = Pt(13)
doc.styles["Heading 1"].font.bold = True
doc.styles["Heading 1"].paragraph_format.space_before = Pt(16)
doc.styles["Heading 1"].paragraph_format.space_after = Pt(7)
doc.core_properties.title = "Portable business evidence for AI recommendations"
doc.core_properties.author = "Joey | Black 4"
doc.core_properties.subject = "Invitation to systems of record and AI platforms"
blocks = (root / "docs/CEO_BRIEF.md").read_text().strip().split("\n\n")
for block in blocks:
    if block.startswith("# "):
        doc.add_paragraph(block[2:], "Title")
    elif block.startswith("## "):
        heading = block[3:]
        p = doc.add_paragraph(heading, "Heading 1")
        if heading == "Make performance count in recommendations":
            p.paragraph_format.page_break_before = True
    elif block.startswith("Reference and proposal:"):
        p = doc.add_paragraph("Reference and proposal: ")
        url = block.split(": ", 1)[1]
        rel = p.part.relate_to(
            url,
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
            is_external=True,
        )
        h = OxmlElement("w:hyperlink")
        h.set(qn("r:id"), rel)
        run = OxmlElement("w:r")
        prop = OxmlElement("w:rPr")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "000000")
        prop.append(color)
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        prop.append(u)
        run.append(prop)
        t = OxmlElement("w:t")
        t.text = "Open Commerce Performance on GitHub"
        run.append(t)
        h.append(run)
        p._p.append(h)
        for r in p.runs:
            r.font.size = Pt(9)
    else:
        p = doc.add_paragraph()
        for i, part in enumerate(re.split(r"\*\*(.*?)\*\*", block)):
            r = p.add_run(part)
            r.bold = bool(i % 2)
        if block.startswith("Joey ·"):
            p.paragraph_format.space_after = Pt(18)
            for r in p.runs:
                r.font.size = Pt(10)
    # Remove inherited title borders if any.
    for border in list(p._p.xpath("./w:pPr/w:pBdr")) if "p" in locals() else []:
        border.getparent().remove(border)
for element in (doc.styles.element, doc.element):
    for border in list(element.xpath(".//w:pBdr")):
        border.getparent().remove(border)
doc.save(root / "docs/proposal/portable-business-evidence.docx")
print(root / "docs/proposal/portable-business-evidence.docx")
